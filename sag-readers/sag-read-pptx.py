#!/usr/bin/env python3
"""sag-read-pptx.py — default PPTX reader.

Uses python-pptx for slide titles + text frames. Each slide becomes a
`## Slide N` section. Notes are appended under each slide.

If python-pptx isn't installed OR the file is encrypted, we return exit
2 so the exe can try the COM-based ps1 reader.

Implements the SAG reader protocol (see SPEC.md).
"""
from __future__ import annotations
import json
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", line_buffering=True)
    sys.stderr.reconfigure(encoding="utf-8", line_buffering=True)
except Exception:
    pass

EXTENSIONS = [".pptx"]
EXIT_OK = 0
EXIT_UNSUPPORTED = 2


def _meta() -> None:
    json.dump({
        "name": "sag-read-pptx",
        "version": "1.0.0",
        "extensions": EXTENSIONS,
        "languages": ["zh", "en", "ja", "ko", "any"],
        "timeoutMs": 180000,
        "needs": ["python>=3.10", "python-pptx>=0.6"],
        "maxBytes": 524288000,
    }, sys.stdout, ensure_ascii=False)
    sys.stdout.write("\n")


def _probe(path: str) -> int:
    if not os.path.isfile(path):
        return 1
    _, ext = os.path.splitext(path)
    if ext.lower() not in EXTENSIONS:
        return 1
    try:
        with open(path, "rb") as f:
            return 0 if f.read(2) == b"PK" else 1
    except OSError:
        return 1


def _extract(path: str) -> int:
    basename = os.path.basename(path)
    try:
        from pptx import Presentation
    except ImportError:
        sys.stderr.write("python-pptx not installed\n")
        return EXIT_UNSUPPORTED
    try:
        prs = Presentation(path)
    except Exception as e:
        msg = str(e).lower()
        if "encrypt" in msg or "password" in msg or "not a zip" in msg:
            sys.stderr.write(f"encrypted or malformed: {e}\n")
            return EXIT_UNSUPPORTED
        sys.stderr.write(f"read failed: {e}\n")
        return 1

    slide_parts: list[str] = []
    for i, slide in enumerate(prs.slides):
        chunks: list[str] = []
        title = None
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if not text:
                    continue
                if title is None and shape.has_text_frame and shape == slide.shapes.title:
                    title = text
                else:
                    chunks.append(text)
        section_lines: list[str] = []
        if title:
            section_lines.append(f"### {title}")
        for c in chunks:
            section_lines.append(f"- {c}")
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                section_lines.append("")
                section_lines.append(f"_备注_: {notes_text}")
        if section_lines:
            slide_parts.append(f"## Slide {i + 1}\n\n" + "\n".join(section_lines))

    if not slide_parts:
        sys.stderr.write("pptx has no extractable text\n")
        return EXIT_UNSUPPORTED

    body = "\n\n".join(slide_parts)
    out = [f"# {basename}", "", body]
    sys.stdout.write("\n".join(out))
    if not body.endswith("\n"):
        sys.stdout.write("\n")
    sys.stderr.write(json.dumps({
        "title": basename,
        "bytes": os.path.getsize(path),
        "slides": len(prs.slides),
        "extractedAt": _now_iso(),
    }, ensure_ascii=False) + "\n")
    return EXIT_OK


def _now_iso() -> str:
    from datetime import datetime
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _dispatch(argv: list[str]) -> int:
    if len(argv) >= 2 and argv[1] == "--meta":
        _meta()
        return EXIT_OK
    if len(argv) >= 3 and argv[1] == "--probe":
        return _probe(argv[2])
    if len(argv) >= 2:
        return _extract(argv[1])
    sys.stderr.write("usage: sag-read-pptx.py [--meta|--probe F] <file>\n")
    return 1


if __name__ == "__main__":
    sys.exit(_dispatch(sys.argv))